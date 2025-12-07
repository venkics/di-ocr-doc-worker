#!/usr/bin/env python3
"""
di_pages_merge.py — parallel DI per page + append-only final workbook + doc-level KV/Entities splice
- Stage A: prebuilt-layout per-page (bounded parallelism)
- Stage B: per-page export (graph + extract) in parallel
- Stage C: append all page workbooks into <base>_layout_graph.xlsx
- Stage D: ONE doc-level prebuilt-document; add KV_Nodes, EntitiesExtracted, EntitiesMerged,
           and KV-related Edges into the final workbook
"""

import argparse, json, os, sys, time, threading, queue
from pathlib import Path
from typing import Any, Dict, List, Tuple
from concurrent.futures import ThreadPoolExecutor, as_completed
import requests
from pypdf import PdfReader
from openpyxl import Workbook, load_workbook

# project imports
from llm_doc_pipeline.orchestrator.ocr.graph_builder import build_layout_graph
from llm_doc_pipeline.orchestrator.ocr.entities_builder import build_entities
from llm_doc_pipeline.orchestrator.ocr.graph_to_excel import export_graph_to_excel
from llm_doc_pipeline.orchestrator.ocr.excel_export import export_layout_to_excel
from llm_doc_pipeline.orchestrator.ocr.checkbox_linker import link_checkboxes_right_labels
from llm_doc_pipeline.orchestrator.ocr.di_client import require_env, analyze_model

API_VER          = os.getenv("DI_API_VERSION", "2023-07-31")
DI_ENDPOINT      = (os.getenv("AZ_DI_ENDPOINT") or os.getenv("DI_ENDPOINT") or "").rstrip("/")
DI_KEY           = (os.getenv("AZ_DI_KEY") or os.getenv("DI_KEY") or "")
PAGE_PARALLEL    = int(os.getenv("PAGE_PARALLEL", "2"))   # per-page exporters
DI_MAX_INFLIGHT  = int(os.getenv("DI_MAX_INFLIGHT", "1")) # concurrent DI submits
FIRST_PAGE_BURST = int(os.getenv("FIRST_PAGE_BURST", "1"))

if not DI_ENDPOINT or not DI_KEY:
    print("ERROR: DI endpoint/key missing.", file=sys.stderr)
    sys.exit(2)

# ---------- DI helpers ----------
def _analyze_url(model: str, pages: List[int], use_new_route: bool) -> str:
    pages_q = ",".join(str(p) for p in pages)
    route = "documentintelligence" if use_new_route else "formrecognizer"
    return f"{DI_ENDPOINT}/{route}/documentModels/{model}:analyze?api-version={API_VER}&pages={pages_q}"

import mimetypes

def _post_analyze_single_page(file_bytes: bytes, page: int, mime_type: str = "application/pdf") -> Dict[str, Any]:
    headers = {"Ocp-Apim-Subscription-Key": DI_KEY, "Content-Type": mime_type}
    last_txt = ""
    for use_new in (False, True):
        url = _analyze_url("prebuilt-layout", [page], use_new)
        r = requests.post(url, headers=headers, data=file_bytes, timeout=120)
        if r.status_code in (400, 404) and "documentModels" in r.text:
            last_txt = r.text
            continue
        if r.status_code == 429:
            ra = r.headers.get("Retry-After") or r.headers.get("retry-after")
            time.sleep(float(ra) if ra else 2.0)
            r = requests.post(url, headers=headers, data=file_bytes, timeout=120)
        r.raise_for_status()

        op = r.headers.get("operation-location") or r.headers.get("Operation-Location")
        if not op:
            raise RuntimeError(f"No operation-location; body={r.text[:300]}")

        for _ in range(300):
            pr = requests.get(op, headers={"Ocp-Apim-Subscription-Key": DI_KEY}, timeout=30)
            if pr.status_code in (429,) or pr.status_code >= 500:
                ra = pr.headers.get("Retry-After") or pr.headers.get("retry-after")
                time.sleep(float(ra) if ra else 0.5); continue
            pr.raise_for_status()
            data = pr.json()
            st = data.get("status")
            if st in ("succeeded", "partiallySucceeded"):
                return data
            if st == "failed":
                raise RuntimeError(f"DI analyze failed: {data}")
            ra = pr.headers.get("Retry-After") or pr.headers.get("retry-after")
            time.sleep(float(ra) if ra else 0.25)
        raise RuntimeError("Timeout waiting for DI result")
    raise RuntimeError(f"Route mismatch; last_resp={last_txt[:300]}")

def _save_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")

# ---------- Pipeline (pages) ----------
def run_pipelined(pdf_path: Path, out_dir: Path) -> List[int]:
    file_bytes = pdf_path.read_bytes()
    
    # Detect mime type
    mime_type, _ = mimetypes.guess_type(pdf_path)
    if not mime_type:
        # Fallback based on extension
        ext = pdf_path.suffix.lower()
        if ext in (".jpg", ".jpeg"): mime_type = "image/jpeg"
        elif ext == ".png": mime_type = "image/png"
        elif ext == ".tiff": mime_type = "image/tiff"
        elif ext == ".bmp": mime_type = "image/bmp"
        else: mime_type = "application/pdf"

    if mime_type == "application/pdf":
        try:
            total_pages = len(PdfReader(str(pdf_path)).pages)
        except Exception:
            print(f"WARNING: pypdf failed on {pdf_path}, assuming 1 page", file=sys.stderr)
            total_pages = 1
    else:
        total_pages = 1

    order = list(range(1, total_pages + 1))
    if FIRST_PAGE_BURST > 0 and total_pages > 1:
        order = order[:FIRST_PAGE_BURST] + order[FIRST_PAGE_BURST:]

    base   = pdf_path.stem
    gp_dir = out_dir / "graph_pages"
    gp_dir.mkdir(parents=True, exist_ok=True)

    q: "queue.Queue[Tuple[int, Dict[str, Any]]]" = queue.Queue()
    errors: List[Tuple[int, str]] = []
    done_pages: List[int] = []
    lock = threading.Lock()
    t0_doc = time.time()

    def producer():
        try:
            with ThreadPoolExecutor(max_workers=max(1, DI_MAX_INFLIGHT)) as ex:
                futs = {ex.submit(_post_analyze_single_page, file_bytes, pn, mime_type): pn for pn in order}
                for fut in as_completed(futs):
                    pn = futs[fut]
                    data = fut.result()
                    wrapped = {
                        "apiVersion": API_VER,
                        "modelId": "prebuilt-layout",
                        "contentFormat": "text",
                        "stringIndexType": "textElements",
                        "analyzeResult": data.get("analyzeResult") or data,
                    }
                    q.put((pn, wrapped))
                    print(f"[DI] page {pn} ready", flush=True)
        except Exception as e:
            errors.append((-1, repr(e)))
        finally:
            q.put((-1, {}))  # sentinel

    def consumer(idx: int):
        while True:
            pn, wrapped = q.get()
            if pn == -1:
                q.put((-1, {})); return
            t0 = time.time()
            try:
                graph = build_layout_graph(wrapped, None)
                ents  = build_entities(None)  # none per-page; doc-level later
                cbs   = link_checkboxes_right_labels(wrapped, only_selected=True)

                _save_json(gp_dir / f"{base}_layout.p{pn:04d}.json", wrapped)
                _save_json(gp_dir / f"{base}_document.p{pn:04d}.json", wrapped)
                _save_json(gp_dir / f"{base}_layout_graph.p{pn:04d}.json", graph)
                _save_json(gp_dir / f"{base}_entities_extracted.p{pn:04d}.json", ents)
                _save_json(gp_dir / f"{base}_checkboxes_extracted.p{pn:04d}.json", cbs)

                out_graph   = gp_dir / f"{base}_layout_graph.p{pn:04d}.xlsx"
                out_extract = gp_dir / f"{base}_layout_extract.p{pn:04d}.xlsx"
                export_graph_to_excel(wrapped, graph, ents, str(out_graph))
                export_layout_to_excel(wrapped, str(out_extract))
                if not out_graph.exists():
                    raise RuntimeError(f"Exporter did not write {out_graph.name}")

                with lock:
                    done_pages.append(pn)
                print(f"[EXPORT] page {pn} in {time.time()-t0:.2f}s (worker {idx})", flush=True)
            except Exception as e:
                errors.append((pn, repr(e)))
                print(f"[EXPORT] page {pn} ERROR: {e}", flush=True)

    threading.Thread(target=producer, daemon=True).start()
    consumers = [threading.Thread(target=consumer, args=(i,), daemon=True)
                 for i in range(max(1, PAGE_PARALLEL))]
    for t in consumers: t.start()
    for t in consumers: t.join()

    if errors:
        pn, err = errors[0]
        raise RuntimeError(f"Page {pn} failed: {err}")

    done_pages.sort()
    print(f"[DOC] all pages done in {time.time() - t0_doc:.2f}s", flush=True)
    return done_pages

# ---------- Final workbook (append per-page) ----------
def _headers_from(book_path: Path, sheets: List[str]) -> Dict[str, List[Any]]:
    headers: Dict[str, List[Any]] = {}
    wb = load_workbook(book_path, data_only=True)
    try:
        for s in sheets:
            if s in wb.sheetnames:
                headers[s] = [c.value for c in next(wb[s].iter_rows(min_row=1, max_row=1))]
    finally:
        wb.close()
    return headers

def _append_rows(dst_ws, src_ws, start_row: int = 2):
    for row in src_ws.iter_rows(min_row=start_row, values_only=True):
        dst_ws.append(row)

def _sheetnames(book_path: Path) -> List[str]:
    wb = load_workbook(book_path, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()

def build_final_workbook(pdf_path: Path, out_dir: Path, pages: List[int]) -> Path:
    base   = pdf_path.stem
    gp_dir = out_dir / "graph_pages"
    final  = out_dir / f"{base}_layout_graph.xlsx"

    def first_match(kind: str) -> Path | None:
        for pn in pages:
            p = gp_dir / f"{base}_{kind}.p{pn:04d}.xlsx"
            if p.exists(): return p
        return None

    first_graph   = first_match("layout_graph")
    first_extract = first_match("layout_extract")
    if not first_graph:
        raise FileNotFoundError("No per-page graph XLSX found to merge.")

    desired_graph = ["LayoutGraph", "Tables", "Cells", "KV_Nodes", "Edges", "EntitiesExtracted", "EntitiesMerged"]
    g_names = _sheetnames(first_graph)
    graph_order = [s for s in desired_graph if s in g_names] + [s for s in g_names if s not in desired_graph]

    extract_order: List[str] = []
    if first_extract:
        desired_extract = ["SelectionMarks", "CheckboxesLabeled"]
        e_names = _sheetnames(first_extract)
        extract_order = [s for s in desired_extract if s in e_names] + [s for s in e_names if s not in desired_extract]

    headers = {}
    headers.update(_headers_from(first_graph, graph_order))
    if first_extract:
        headers.update(_headers_from(first_extract, extract_order))

    final_wb = Workbook()
    final_wb.remove(final_wb.active)
    for s in graph_order:
        ws = final_wb.create_sheet(title=s); ws.append(headers[s])
    for s in extract_order:
        if s in headers:
            ws = final_wb.create_sheet(title=s); ws.append(headers[s])

    for pn in pages:
        g_path = gp_dir / f"{base}_layout_graph.p{pn:04d}.xlsx"
        if g_path.exists():
            gwb = load_workbook(g_path, data_only=True)
            try:
                for s in graph_order:
                    if s in gwb.sheetnames and s in final_wb.sheetnames:
                        _append_rows(final_wb[s], gwb[s], start_row=2)
            finally:
                gwb.close()

        e_path = gp_dir / f"{base}_layout_extract.p{pn:04d}.xlsx"
        if e_path.exists() and extract_order:
            ewb = load_workbook(e_path, data_only=True)
            try:
                for s in extract_order:
                    if s in ewb.sheetnames and s in final_wb.sheetnames:
                        _append_rows(final_wb[s], ewb[s], start_row=2)
            finally:
                ewb.close()

    final_wb.save(final)
    print(f"[EXCEL] wrote {final}", flush=True)
    return final

# ---------- Doc-level KV/Entities splice ----------
def _combine_layout_from_pages(out_dir: Path, base: str, pages: List[int]) -> Dict[str, Any]:
    """Build one layout_json by concatenating the saved per-page layout JSONs."""
    gp_dir = out_dir / "graph_pages"
    pages_list, tables_list = [], []
    for pn in pages:
        p_json = (gp_dir / f"{base}_layout.p{pn:04d}.json").read_text(encoding="utf-8")
        wrapped = json.loads(p_json)
        ar = wrapped.get("analyzeResult", {}) or {}
        if ar.get("pages"):
            pages_list.extend(ar["pages"])
        if ar.get("tables"):
            tables_list.extend(ar["tables"])
    return {
        "apiVersion": API_VER,
        "modelId": "prebuilt-layout",
        "contentFormat": "text",
        "stringIndexType": "textElements",
        "analyzeResult": {"pages": pages_list, "tables": tables_list}
    }

def splice_doclevel_kv_entities(pdf_path: Path, out_dir: Path, pages: List[int], final_xlsx: Path) -> None:
    """Run one prebuilt-document, build KV graph+entities, append sheets to final workbook."""
    base = pdf_path.stem
    # DI credentials (use your helper)
    endpoint, key = require_env()
    # analyze prebuilt-document (1 call)
    document = analyze_model(endpoint, key, "prebuilt-document", str(pdf_path))
    # entities
    entities = build_entities(document)  # {"entities":[...]} with pages + bboxes
    # combined layout (from the per-page saved files) for KV nodes/edges alignment
    layout_combined = _combine_layout_from_pages(out_dir, base, pages)
    # build graph with KV nodes/edges
    graph_doc = build_layout_graph(layout_combined, document)

    # write a small temp workbook containing just the doc-level sets we want,
    # then copy those sheets into the final workbook.
    tmp = out_dir / f"{base}__doclevel_kv.xlsx"
    export_graph_to_excel(layout_combined, graph_doc, entities, str(tmp))

    # Copy sheets: KV_Nodes, EntitiesExtracted, EntitiesMerged, and append KV edges
    kv_sheets = ["KV_Nodes", "EntitiesExtracted", "EntitiesMerged"]
    src = load_workbook(tmp, data_only=True)
    dst = load_workbook(final_xlsx, data_only=False)
    try:
        # replace/create KV/Entities sheets
        for s in kv_sheets:
            if s in dst.sheetnames:
                del dst[s]
            if s in src.sheetnames:
                dst.create_sheet(s)
                dws = dst[s]; sws = src[s]
                # write header
                headers = [c.value for c in next(sws.iter_rows(min_row=1, max_row=1))]
                dws.append(headers)
                # append rows
                for row in sws.iter_rows(min_row=2, values_only=True):
                    dws.append(row)
        # append KV-related edges into existing Edges sheet (create if missing)
        if "Edges" not in dst.sheetnames:
            dst.create_sheet("Edges").append(["src","dst","rel"])
        ews = dst["Edges"]
        if "Edges" in src.sheetnames:
            for row in src["Edges"].iter_rows(min_row=2, values_only=True):
                # keep only kv-pair / overlaps edges (avoid duplicating word edges)
                if row and row[2] in ("kv-pair","overlaps"):
                    ews.append(row)
        dst.save(final_xlsx)
    finally:
        src.close()
        dst.close()
    try:
        tmp.unlink()
    except Exception:
        pass
    print(f"[EXCEL] spliced doc-level KV/Entities into {final_xlsx}", flush=True)

def write_manifests(pdf_path: Path, out_dir: Path, pages: List[int]) -> None:
    base = pdf_path.stem
    manifest = {
        "type": "graph_pages_manifest",
        "apiVersion": API_VER,
        "pages": len(pages),
        "graph_pages_dir": "graph_pages",
        "base": base,
        "pageFiles": [f"{base}_layout_graph.p{pn:04d}.xlsx" for pn in pages],
        "pageExtractFiles": [f"{base}_layout_extract.p{pn:04d}.xlsx" for pn in pages],
    }
    (out_dir / "layout.json").write_text(json.dumps(manifest, indent=2), "utf-8")
    (out_dir / "document.json").write_text(json.dumps(manifest, indent=2), "utf-8")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out_dir", required=True)
    args = ap.parse_args()

    pdf = Path(args.input).resolve()
    out = Path(args.out_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)

    pages = run_pipelined(pdf, out)
    final_xlsx = build_final_workbook(pdf, out, pages)
    # NEW: one doc-level DI for KV/Entities; splice into final workbook
    splice_doclevel_kv_entities(pdf, out, pages, final_xlsx)
    write_manifests(pdf, out, pages)

if __name__ == "__main__":
    main()
